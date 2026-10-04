"""多格式内容提取：PDF / Word / PPT / Excel / 纯文本 / 图片渲染。"""
from __future__ import annotations

import csv
import io
import zipfile
import zlib
from pathlib import Path

from . import config


class UnreadableFile(Exception):
    """源文件本身读不出来：空的 / 不是这个格式 / 内部结构损坏。

    跟「程序出错」区分开：这类文件重试多少次结果都一样，也绝不允许去改动源文件，
    所以索引器把它们标成「无法解析」并把原因记进 text_note / vision_note，
    而不是混在「失败」里让人反复点重试。
    """


_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_OOXML_LABELS = {"word": "Word 文档（.docx）", "ppt": "PowerPoint 演示文稿（.pptx）",
                 "excel": "Excel 表格（.xlsx）"}


def _ooxml_errors() -> tuple:
    """第三方库遇到"包里少东西 / 压缩流坏了"时抛的几种异常，全归到「无法解析」。"""
    errs = [zipfile.BadZipFile, zlib.error]
    try:
        from docx.opc.exceptions import PackageNotFoundError as _docx_err
        errs.append(_docx_err)
    except Exception:
        pass
    try:
        from pptx.exc import PackageNotFoundError as _pptx_err
        errs.append(_pptx_err)
    except Exception:
        pass
    try:
        from openpyxl.utils.exceptions import InvalidFileException as _xlsx_err
        errs.append(_xlsx_err)
    except Exception:
        pass
    return tuple(errs)


BAD_PACKAGE_ERRORS = _ooxml_errors()


def _precheck(path: Path, kind: str, ext: str) -> None:
    """交给第三方库之前先看一眼文件本身，这样报出来的原因是人话。"""
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise UnreadableFile("读不到这个文件：" + str(exc))
    if size <= 0:
        raise UnreadableFile("这个文件是空的（0 字节），没有内容可以提取。")
    try:
        with path.open("rb") as fh:
            head = fh.read(8)
    except OSError as exc:
        raise UnreadableFile("读不到这个文件：" + str(exc))
    if kind in _OOXML_LABELS:
        label = _OOXML_LABELS[kind]
        if head.startswith(_OLE2_MAGIC):
            raise UnreadableFile(
                "这是老式 Office 文件（Word / PPT 97-2003 格式），只是后缀被改成了 ." + ext
                + "。程序只读不改，所以请你自己用 Office / WPS 打开它，"
                "「另存为」成 ." + ext + " 之后再重新扫描一次。")
        if not head.startswith(b"PK"):
            raise UnreadableFile("这个文件的内容不是有效的" + label + "（文件可能已损坏，或者没下载/复制完整）。")

IMAGE_EXTS = {"png", "jpg", "jpeg", "gif", "bmp", "webp", "tif", "tiff", "heic", "avif"}
VIDEO_EXTS = {"mp4", "mpg", "mpeg", "rm", "rmvb", "avi", "mkv", "mov", "wmv", "flv", "m4v", "3gp", "ts", "swf"}
AUDIO_EXTS = {"mp3", "wav", "m4a", "flac", "aac", "ogg", "wma"}
PDF_EXTS = {"pdf"}
WORD_EXTS = {"doc", "docx", "rtf"}
PPT_EXTS = {"ppt", "pptx"}
EXCEL_EXTS = {"xls", "xlsx", "csv"}
TEXT_EXTS = {"txt", "md", "markdown", "json", "xml", "html", "htm", "log"}
ARCHIVE_EXTS = {"zip", "rar", "7z", "tar", "gz", "iso"}
WEB_EXTS = {"html", "htm", "mhtml", "webarchive"}
CODE_EXTS = {"py", "js", "c", "cpp", "java", "cs", "html", "css"}

TEXT_EXTRACTABLE = PDF_EXTS | WORD_EXTS | PPT_EXTS | EXCEL_EXTS | TEXT_EXTS

KIND_LABELS = {
    "pdf": "PDF",
    "word": "文档",
    "ppt": "演示文稿",
    "excel": "表格",
    "image": "图片",
    "video": "视频",
    "audio": "音频",
    "text": "文本",
    "web": "网页",
    "archive": "压缩包",
    "other": "其他",
}


def ext_of(name: str) -> str:
    return Path(name).suffix.lower().lstrip(".")


def kind_of(ext: str) -> str:
    ext = (ext or "").lower()
    if ext in PDF_EXTS:
        return "pdf"
    if ext in WORD_EXTS:
        return "word"
    if ext in PPT_EXTS:
        return "ppt"
    if ext in EXCEL_EXTS:
        return "excel"
    if ext in IMAGE_EXTS:
        return "image"
    if ext in VIDEO_EXTS:
        return "video"
    if ext in AUDIO_EXTS:
        return "audio"
    if ext in WEB_EXTS:
        return "web"
    if ext in TEXT_EXTS or ext in CODE_EXTS:
        return "text"
    if ext in ARCHIVE_EXTS:
        return "archive"
    return "other"


def decode_bytes(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030", "big5"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "ignore")


def _clip(text: str) -> str:
    return text.strip()


def extract_pdf(path: Path):
    import pymupdf

    pages = []
    try:
        doc = pymupdf.open(str(path))
    except Exception as exc:
        raise UnreadableFile("这个 PDF 打不开（" + type(exc).__name__
                             + "），文件可能已损坏，或者没下载 / 复制完整。")
    try:
        total = doc.page_count
        limit = min(total, config.EXTRACT_MAX_PAGES)
        for i in range(limit):
            try:
                raw = doc[i].get_text("text")
            except Exception:
                raw = ""
            pages.append((i + 1, _clip(raw or "")))
    finally:
        doc.close()
    return pages, total


def extract_docx(path: Path):
    import docx

    doc = docx.Document(str(path))
    parts = []
    for para in doc.paragraphs:
        t = (para.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return [(1, "\n".join(parts))], 1


def extract_pptx(path: Path):
    from pptx import Presentation

    prs = Presentation(str(path))
    out = []
    for idx, slide in enumerate(prs.slides, 1):
        buf = []
        for shape in slide.shapes:
            try:
                if shape.has_text_frame:
                    t = (shape.text_frame.text or "").strip()
                    if t:
                        buf.append(t)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [(c.text or "").strip() for c in row.cells]
                        if any(cells):
                            buf.append(" | ".join(cells))
            except Exception:
                continue
        try:
            if slide.has_notes_slide:
                note = (slide.notes_slide.notes_text_frame.text or "").strip()
                if note:
                    buf.append("[备注] " + note)
        except Exception:
            pass
        out.append((idx, "\n".join(buf)))
    return out, len(out)


def extract_xlsx(path: Path):
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out = []
    try:
        for idx, ws in enumerate(wb.worksheets, 1):
            lines = ["# 工作表: " + str(ws.title)]
            rows = ws.iter_rows(
                max_row=config.XLSX_MAX_ROWS,
                max_col=config.XLSX_MAX_COLS,
                values_only=True,
            )
            for row in rows:
                cells = ["" if v is None else str(v) for v in row]
                line = " | ".join(cells).strip(" |")
                if line:
                    lines.append(line)
            out.append((idx, "\n".join(lines)))
    finally:
        wb.close()
    return out, len(out)


def extract_csv(path: Path):
    raw = path.read_bytes()[: config.TXT_MAX_BYTES]
    text = decode_bytes(raw)
    lines = []
    for row in csv.reader(io.StringIO(text)):
        line = " | ".join(x.strip() for x in row).strip(" |")
        if line:
            lines.append(line)
    return [(1, "\n".join(lines))], 1


def extract_textfile(path: Path):
    raw = path.read_bytes()[: config.TXT_MAX_BYTES]
    return [(1, decode_bytes(raw).strip())], 1


SKIPPED_EXTS = {"doc", "ppt", "xls"}


def extract(path: Path, ext: str):
    """返回 (pages, total_pages, note)。pages 为 [(页码, 文本)]。

    源文件本身读不出来时抛 UnreadableFile（带中文原因），由索引器标成「无法解析」。
    """
    ext = (ext or "").lower()
    if ext in ("pdf", "docx", "pptx", "xlsx"):
        _precheck(path, kind_of(ext), ext)
    try:
        if ext == "pdf":
            pages, total = extract_pdf(path)
            return pages, total, ""
        if ext == "docx":
            pages, total = extract_docx(path)
            return pages, total, ""
        if ext == "pptx":
            pages, total = extract_pptx(path)
            return pages, total, ""
        if ext == "xlsx":
            pages, total = extract_xlsx(path)
            return pages, total, ""
    except UnreadableFile:
        raise
    except BAD_PACKAGE_ERRORS as exc:
        raise UnreadableFile("这个文件的内部结构已损坏，读不出内容（" + type(exc).__name__
                             + "）。文件可能是在复制 / 下载过程中断掉的，"
                             "可以重新从原处拷一份再扫描。")
    if ext == "csv":
        pages, total = extract_csv(path)
        return pages, total, ""
    if ext in TEXT_EXTS:
        pages, total = extract_textfile(path)
        return pages, total, ""
    if ext in SKIPPED_EXTS:
        return [], 0, "旧版二进制格式（" + ext + "），需先另存为新格式才能提取文字"
    return [], 0, "不支持的格式"


def render_pdf_page(path: Path, page_index: int, max_side: int = None, quality: int = None) -> bytes:
    import pymupdf

    max_side = max_side or config.RENDER_MAX_SIDE
    quality = quality or config.RENDER_QUALITY
    doc = pymupdf.open(str(path))
    try:
        if page_index < 0 or page_index >= doc.page_count:
            raise IndexError("page out of range")
        page = doc[page_index]
        rect = page.rect
        longest = max(rect.width, rect.height) or 1.0
        zoom = min(3.0, max(0.2, float(max_side) / float(longest)))
        pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
        return pix.tobytes("jpeg", jpg_quality=quality)
    finally:
        doc.close()


def image_to_jpeg_bytes(path: Path, max_side: int = None, quality: int = None) -> bytes:
    from PIL import Image

    max_side = max_side or config.RENDER_MAX_SIDE
    quality = quality or config.RENDER_QUALITY
    with Image.open(str(path)) as im:
        im = im.convert("RGB")
        im.thumbnail((max_side, max_side))
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=quality)
        return buf.getvalue()


def pdf_page_count(path: Path) -> int:
    import pymupdf

    doc = pymupdf.open(str(path))
    try:
        return doc.page_count
    finally:
        doc.close()
